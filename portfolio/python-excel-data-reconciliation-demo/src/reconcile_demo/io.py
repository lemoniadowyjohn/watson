from __future__ import annotations

import csv
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .reconciliation import ReconciliationRow


def read_records(path: str | Path) -> list[dict[str, object]]:
    path = Path(path)
    suffix = path.suffix.casefold()
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            return [dict(row) for row in csv.DictReader(handle)]
    if suffix in {".xlsx", ".xlsm"}:
        wb = load_workbook(path, data_only=True, read_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(value).strip() if value is not None else "" for value in rows[0]]
        return [
            dict(zip(headers, row))
            for row in rows[1:]
            if any(value is not None for value in row)
        ]
    raise ValueError(f"Unsupported input format: {path.suffix}")


def write_report(rows: list[ReconciliationRow], output: str | Path) -> None:
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "results"
    headers = list(rows[0].as_dict()) if rows else [
        "component_id",
        "variant",
        "reconciliation_status",
        "reason_codes",
        "description_a",
        "description_b",
        "quantity_a",
        "quantity_b",
        "status_a",
        "status_b",
    ]
    ws.append(headers)
    for row in rows:
        data = row.as_dict()
        ws.append([data[header] for header in headers])

    exceptions = wb.create_sheet("exceptions")
    exceptions.append(headers)
    for row in rows:
        if row.reconciliation_status != "MATCH":
            data = row.as_dict()
            exceptions.append([data[header] for header in headers])

    summary = wb.create_sheet("summary")
    summary.append(["status", "count"])
    counts = Counter(row.reconciliation_status for row in rows)
    for status in ("MATCH", "MISMATCH", "ONLY_IN_A", "ONLY_IN_B", "INVALID"):
        summary.append([status, counts.get(status, 0)])
    wb.save(output)
